"""
Document extractor module for KnowFlow AI.
Supports PDF, TXT, and DOCX document extraction with metadata preservation.
"""

import io
import logging
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Union

import docx
import pymupdf

from src.utils.helpers import clean_filename

logger = logging.getLogger(__name__)


class DocumentExtractionError(Exception):
    """Custom exception for document extraction failures."""
    pass


def extract_from_pdf(
    source_bytes: bytes, filename: str
) -> List[Dict[str, Any]]:
    """
    Extract text page-by-page from a PDF document.
    
    Returns:
        List of dicts: [{"text": str, "source": str, "page": int}, ...]
    """
    documents: List[Dict[str, Any]] = []
    try:
        # Open PDF from bytes in memory
        with pymupdf.open(stream=source_bytes, filetype="pdf") as pdf_doc:
            if pdf_doc.page_count == 0:
                logger.warning("PDF '%s' has 0 pages.", filename)
                return []

            for page_idx, page in enumerate(pdf_doc, start=1):
                page_text = page.get_text() or ""
                if page_text.strip():
                    documents.append({
                        "text": page_text,
                        "source": filename,
                        "page": page_idx,
                    })
    except Exception as e:
        logger.error("Failed to extract PDF '%s': %s", filename, e)
        raise DocumentExtractionError(f"Error parsing PDF '{filename}': {str(e)}") from e

    return documents


def extract_from_txt(
    source_bytes: bytes, filename: str
) -> List[Dict[str, Any]]:
    """
    Extract text from a plain text (.txt) document.
    
    Returns:
        List of dicts: [{"text": str, "source": str, "page": None}]
    """
    try:
        # Attempt standard utf-8 first, fallback to common encodings
        try:
            text = source_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = source_bytes.decode("latin-1")
            except UnicodeDecodeError:
                text = source_bytes.decode("cp1252", errors="replace")

        if not text.strip():
            return []

        return [{
            "text": text,
            "source": filename,
            "page": None,
        }]
    except Exception as e:
        logger.error("Failed to extract TXT '%s': %s", filename, e)
        raise DocumentExtractionError(f"Error reading text file '{filename}': {str(e)}") from e


def extract_from_docx(
    source_bytes: bytes, filename: str
) -> List[Dict[str, Any]]:
    """
    Extract text from a Word document (.docx) including paragraphs and tables.
    
    Returns:
        List of dicts: [{"text": str, "source": str, "page": None}]
    """
    try:
        doc_stream = io.BytesIO(source_bytes)
        doc = docx.Document(doc_stream)
        
        extracted_sections: List[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            content = para.text.strip()
            if content:
                extracted_sections.append(content)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate repeated cell texts from merged cells
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    extracted_sections.append(" | ".join(unique_cells))

        full_text = "\n\n".join(extracted_sections).strip()
        if not full_text:
            return []

        return [{
            "text": full_text,
            "source": filename,
            "page": None,
        }]
    except Exception as e:
        logger.error("Failed to extract DOCX '%s': %s", filename, e)
        raise DocumentExtractionError(f"Error parsing Word file '{filename}': {str(e)}") from e


def extract_documents(
    file_input: Union[str, Path, bytes, BinaryIO, Any],
    filename: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """
    Primary ingestion dispatcher:
    1. Detects file type (PDF, TXT, DOCX).
    2. Extracts text and preserves metadata (e.g., page numbers for PDF).
    3. Handles extraction errors gracefully without crashing the application.
    
    Args:
        file_input: File path, bytes, or file-like object (e.g. Streamlit UploadedFile)
        filename: Optional explicit filename. If omitted, inferred from file_input.
        
    Returns:
        List of structured document items:
        [{"text": "...", "source": "file.pdf", "page": 1}, ...]
    """
    # 1. Resolve raw bytes and filename
    raw_bytes: bytes
    resolved_name: str = filename or "document"

    if isinstance(file_input, (str, Path)):
        path_obj = Path(file_input)
        resolved_name = filename or path_obj.name
        if not path_obj.exists():
            logger.error("File does not exist: %s", file_input)
            return []
        with open(path_obj, "rb") as f:
            raw_bytes = f.read()
    elif hasattr(file_input, "read"):
        # File-like object (e.g. Streamlit UploadedFile)
        if hasattr(file_input, "name") and not filename:
            resolved_name = file_input.name
        if hasattr(file_input, "getvalue"):
            raw_bytes = file_input.getvalue()
        else:
            raw_bytes = file_input.read()
            if hasattr(file_input, "seek"):
                file_input.seek(0)
    elif isinstance(file_input, (bytes, bytearray)):
        raw_bytes = bytes(file_input)
    else:
        logger.error("Unsupported file_input type: %s", type(file_input))
        return []

    safe_name = clean_filename(resolved_name)
    ext = Path(safe_name).suffix.lower()

    # 2. Dispatch based on detected extension
    try:
        if ext == ".pdf":
            return extract_from_pdf(raw_bytes, safe_name)
        elif ext == ".txt":
            return extract_from_txt(raw_bytes, safe_name)
        elif ext in (".docx", ".doc"):
            if ext == ".doc":
                logger.warning("Legacy .doc format detected for '%s'. Attempting docx parser.", safe_name)
            return extract_from_docx(raw_bytes, safe_name)
        else:
            logger.warning("Unsupported file format '%s' for file '%s'.", ext, safe_name)
            return []
    except DocumentExtractionError as de:
        logger.warning("Gracefully handled extraction error: %s", de)
        return []
    except Exception as e:
        logger.error("Unexpected error extracting '%s': %s", safe_name, e)
        return []
