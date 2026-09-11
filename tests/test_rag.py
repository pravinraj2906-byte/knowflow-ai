"""
Unit and integration tests for KnowFlow AI RAG pipeline.
Covers text chunking, metadata preservation, file parsing, stable IDs,
empty document handling, and hallucination control.
All external Gemini API calls are mocked to run offline without an API key.
"""

import io
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

# Ensure root directory is on Python path
ROOT_DIR = Path(__file__).parent.parent.resolve()
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import docx
import pymupdf
import pytest

from src.ingestion.extractor import extract_documents
from src.ingestion.chunker import (
    chunk_documents,
    chunk_text,
    split_text_with_overlap,
)
from src.rag.retriever import Retriever, INSUFFICIENT_CONTEXT_MESSAGE
from src.rag.generator import GroundedGenerator
from src.rag.vector_store import VectorStore
from src.utils.helpers import (
    generate_chunk_id,
    sanitize_text,
    format_source_citation,
)


# =============================================================================
# 1. TEXT CHUNKING TESTS
# =============================================================================

def test_split_text_with_overlap_boundaries():
    """Verify that split_text_with_overlap splits long text with overlap and respects boundaries."""
    paragraph1 = "First paragraph containing important introductory information about KnowFlow AI."
    paragraph2 = "Second paragraph explaining the retrieval-augmented generation architecture in detail."
    paragraph3 = "Third paragraph detailing how ChromaDB stores document vectors and metadata persistently."
    combined = f"{paragraph1}\n\n{paragraph2}\n\n{paragraph3}"

    chunks = split_text_with_overlap(
        combined,
        chunk_size=120,
        chunk_overlap=30,
        min_chunk_size=20,
    )

    assert len(chunks) >= 2
    for chunk in chunks:
        assert len(chunk) >= 20
        assert isinstance(chunk, str)


def test_split_text_short_text():
    """Short text smaller than chunk_size should return as a single chunk."""
    short_text = "This is a brief piece of text."
    chunks = split_text_with_overlap(short_text, chunk_size=500, min_chunk_size=10)
    assert len(chunks) == 1
    assert chunks[0] == short_text


def test_split_text_tiny_chunks_filtered():
    """Trivial chunks smaller than min_chunk_size should be omitted."""
    tiny_text = "Hi."
    chunks = split_text_with_overlap(tiny_text, chunk_size=500, min_chunk_size=10)
    assert len(chunks) == 0


# =============================================================================
# 2. METADATA PRESERVATION TESTS
# =============================================================================

def test_metadata_preservation_across_chunks():
    """Verify that metadata (source, page, chunk_index) is preserved on every chunk."""
    extracted_docs = [
        {"text": "Content from page one of the guidelines.", "source": "guide.pdf", "page": 1},
        {"text": "Content from page two of the guidelines.", "source": "guide.pdf", "page": 2},
        {"text": "Notes from the developer memo.", "source": "notes.txt", "page": None},
    ]

    chunks = chunk_documents(extracted_docs, chunk_size=200, min_chunk_size=10)

    assert len(chunks) == 3
    # Check first chunk
    assert chunks[0]["source"] == "guide.pdf"
    assert chunks[0]["page"] == 1
    assert chunks[0]["chunk_index"] == 0

    # Check second chunk
    assert chunks[1]["source"] == "guide.pdf"
    assert chunks[1]["page"] == 2
    assert chunks[1]["chunk_index"] == 1

    # Check third chunk (TXT without page)
    assert chunks[2]["source"] == "notes.txt"
    assert chunks[2]["page"] is None
    assert chunks[2]["chunk_index"] == 2


def test_format_source_citation():
    """Test source citation formatting rules."""
    pdf_meta = {"source": "handbook.pdf", "page": 4}
    assert format_source_citation(pdf_meta) == "handbook.pdf — Page 4"

    docx_meta = {"source": "report.docx", "page": None}
    assert format_source_citation(docx_meta) == "report.docx"

    txt_meta = {"source": "notes.txt", "page": -1}
    assert format_source_citation(txt_meta) == "notes.txt"


# =============================================================================
# 3. SUPPORTED FILE HANDLING LOGIC TESTS
# =============================================================================

def test_pdf_extraction():
    """Test extraction of multi-page PDF bytes with page numbers preserved."""
    pdf_doc = pymupdf.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "Page 1: Overview of KnowFlow AI system.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "Page 2: Vector search indexing pipeline.")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    results = extract_documents(pdf_bytes, filename="overview.pdf")
    assert len(results) == 2
    assert results[0]["source"] == "overview.pdf"
    assert results[0]["page"] == 1
    assert "Page 1" in results[0]["text"]

    assert results[1]["source"] == "overview.pdf"
    assert results[1]["page"] == 2
    assert "Page 2" in results[1]["text"]


def test_docx_extraction():
    """Test extraction of Word DOCX document with paragraphs and tables."""
    doc = docx.Document()
    doc.add_paragraph("Internship Project Requirements.")
    doc.add_paragraph("Features include ChromaDB and Streamlit.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Framework"
    table.rows[0].cells[1].text = "Google GenAI"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    results = extract_documents(docx_bytes, filename="requirements.docx")
    assert len(results) == 1
    assert results[0]["source"] == "requirements.docx"
    assert results[0]["page"] is None
    assert "Internship Project Requirements" in results[0]["text"]
    assert "Framework | Google GenAI" in results[0]["text"]


def test_txt_extraction():
    """Test extraction of plain text files."""
    content = "Quick meeting notes:\n- Review Week 2 submission\n- Verify tests pass"
    txt_bytes = content.encode("utf-8")

    results = extract_documents(txt_bytes, filename="notes.txt")
    assert len(results) == 1
    assert results[0]["source"] == "notes.txt"
    assert results[0]["page"] is None
    assert "Review Week 2 submission" in results[0]["text"]


def test_unsupported_file_format():
    """Test that unsupported formats are safely handled without raising an exception."""
    results = extract_documents(b"fake image data", filename="picture.png")
    assert results == []


# =============================================================================
# 4. STABLE CHUNK IDS TESTS
# =============================================================================

def test_stable_chunk_ids():
    """Verify that chunk IDs are deterministic across repeated runs with same input."""
    text = "KnowFlow AI delivers grounded answers from document context."
    id1 = generate_chunk_id("doc.pdf", page=1, chunk_idx=0, content=text)
    id2 = generate_chunk_id("doc.pdf", page=1, chunk_idx=0, content=text)
    assert id1 == id2
    assert "doc.pdf_p1_c0" in id1

    # Changing content changes the hash part of the ID
    id_different = generate_chunk_id("doc.pdf", page=1, chunk_idx=0, content="Altered content.")
    assert id1 != id_different


# =============================================================================
# 5. EMPTY DOCUMENT HANDLING TESTS
# =============================================================================

def test_empty_text_and_whitespace_handling():
    """Verify empty or whitespace-only documents produce zero chunks without errors."""
    assert extract_documents(b"", filename="empty.txt") == []
    assert extract_documents(b"   \n\n   \t  ", filename="blank.txt") == []

    # Chunker on empty input
    assert chunk_documents([]) == []
    assert chunk_text("", source="empty.txt") == []
    assert sanitize_text("   \n\t  ") == ""


# =============================================================================
# 6. HALLUCINATION & RELEVANCE FALLBACK TESTS
# =============================================================================

def test_retriever_hallucination_fallback_when_distance_exceeds_threshold():
    """
    When the vector search returns chunks with distance higher than the
    relevance threshold, the retriever flags context as insufficient.
    """
    mock_vector_store = MagicMock(spec=VectorStore)
    mock_vector_store.count.return_value = 5

    # Simulate chunk with high distance (low relevance)
    mock_vector_store.similarity_search.return_value = [
        {
            "chunk_id": "c1",
            "text": "Completely unrelated topic about astrophysics.",
            "metadata": {"source": "astro.txt", "page": None},
            "distance": 0.88,  # > default threshold 0.70
            "similarity": 0.12,
        }
    ]

    mock_embedder = MagicMock()
    mock_embedder.embed_query.return_value = [0.1] * 768

    retriever = Retriever(
        vector_store=mock_vector_store,
        embedder=mock_embedder,
        relevance_threshold=0.70,
    )

    result = retriever.retrieve("What is the project submission deadline?")

    assert result.has_relevant_context is False
    assert len(result.chunks) == 0
    assert result.fallback_message == INSUFFICIENT_CONTEXT_MESSAGE


def test_retriever_preserves_relevant_chunks():
    """Chunks within relevance threshold are retained."""
    mock_vector_store = MagicMock(spec=VectorStore)
    mock_vector_store.count.return_value = 5

    mock_vector_store.similarity_search.return_value = [
        {
            "chunk_id": "c1",
            "text": "The project submission deadline is Friday 5 PM.",
            "metadata": {"source": "schedule.pdf", "page": 2},
            "distance": 0.22,  # <= default threshold 0.70
            "similarity": 0.78,
        }
    ]

    mock_embedder = MagicMock()
    mock_embedder.embed_query.return_value = [0.1] * 768

    retriever = Retriever(
        vector_store=mock_vector_store,
        embedder=mock_embedder,
        relevance_threshold=0.70,
    )

    result = retriever.retrieve("When is the project deadline?")

    assert result.has_relevant_context is True
    assert len(result.chunks) == 1
    assert result.chunks[0]["metadata"]["source"] == "schedule.pdf"
    assert result.chunks[0]["metadata"]["page"] == 2


def test_generator_direct_fallback_when_forced():
    """Generator returns required exact fallback string without calling LLM."""
    generator = GroundedGenerator(api_key="mock_key")
    result = generator.generate_answer(
        question="What is the capital of Mars?",
        retrieved_chunks=[],
        force_fallback=True,
    )

    assert result.answer == INSUFFICIENT_CONTEXT_MESSAGE
    assert result.sources == []
    assert result.is_fallback is True


def test_generator_grounded_answer_with_mock():
    """Generator invokes Gemini with grounded context and produces clean sources."""
    generator = GroundedGenerator(api_key="mock_key")

    mock_client = MagicMock()
    mock_response = MagicMock()
    mock_response.text = "KnowFlow AI uses ChromaDB for persistent vector storage."
    mock_client.models.generate_content.return_value = mock_response

    generator._client = mock_client

    test_chunks = [
        {
            "chunk_id": "c1",
            "text": "KnowFlow AI utilizes ChromaDB for fast local embeddings and similarity search.",
            "metadata": {"source": "architecture.pdf", "page": 4},
        },
        {
            "chunk_id": "c2",
            "text": "Embeddings are generated via Google GenAI SDK.",
            "metadata": {"source": "notes.txt", "page": None},
        },
    ]

    result = generator.generate_answer("How are vectors stored?", test_chunks)

    assert "ChromaDB" in result.answer
    assert result.is_fallback is False
    assert "architecture.pdf — Page 4" in result.sources
    assert "notes.txt" in result.sources


# =============================================================================
# 7. VECTOR STORE PERSISTENCE & DEDUPLICATION TESTS
# =============================================================================

def test_vector_store_add_and_deduplication(tmp_path):
    """Verify that adding documents to the vector store works and replaces duplicates."""
    store = VectorStore(persist_directory=tmp_path / "test_chroma", collection_name="test_col")
    
    chunks = [
        {
            "chunk_id": "chunk_1",
            "text": "First chunk content.",
            "source": "doc1.txt",
            "page": None,
            "chunk_index": 0,
        },
        {
            "chunk_id": "chunk_2",
            "text": "Second chunk content.",
            "source": "doc1.txt",
            "page": None,
            "chunk_index": 1,
        },
    ]
    # Synthetic 4-dimensional embeddings for testing
    embeddings = [[0.1, 0.2, 0.3, 0.4], [0.5, 0.6, 0.7, 0.8]]
    
    added = store.add_chunks(chunks, embeddings)
    assert added == 2
    assert store.count() == 2

    # Re-indexing the same file (same source) should replace older chunks without duplicating
    updated_chunks = [
        {
            "chunk_id": "chunk_3",
            "text": "New updated single chunk for doc1.",
            "source": "doc1.txt",
            "page": None,
            "chunk_index": 0,
        }
    ]
    updated_embeddings = [[0.2, 0.3, 0.4, 0.5]]
    store.add_chunks(updated_chunks, updated_embeddings)

    assert store.count() == 1
    docs = store.list_documents()
    assert len(docs) == 1
    assert docs[0]["source"] == "doc1.txt"
    assert docs[0]["chunk_count"] == 1


def test_vector_store_reset(tmp_path):
    """Verify clearing/resetting the knowledge base."""
    store = VectorStore(persist_directory=tmp_path / "test_reset", collection_name="test_reset_col")
    chunks = [
        {
            "chunk_id": "c1",
            "text": "Some text",
            "source": "sample.pdf",
            "page": 1,
            "chunk_index": 0,
        }
    ]
    store.add_chunks(chunks, [[0.1, 0.2, 0.3, 0.4]])
    assert store.count() == 1

    store.reset_collection()
    assert store.count() == 0
    assert store.list_documents() == []


# =============================================================================
# 8. DEFAULT MODEL CONFIGURATION TESTS
# =============================================================================

def test_default_model_configurations():
    """Verify that updated default models (gemini-embedding-2 and gemini-3.6-flash) are active."""
    from src.rag.embeddings import DEFAULT_EMBEDDING_MODEL, GeminiEmbedder
    from src.rag.generator import DEFAULT_GENERATION_MODEL, GroundedGenerator

    assert DEFAULT_EMBEDDING_MODEL == "gemini-embedding-2"
    assert DEFAULT_GENERATION_MODEL == "gemini-3.6-flash"

    embedder = GeminiEmbedder(api_key="mock_key")
    assert embedder.model_name == "gemini-embedding-2"

    generator = GroundedGenerator(api_key="mock_key")
    assert generator.model_name == "gemini-3.6-flash"


# =============================================================================
# 9. GEMINI-EMBEDDING-2 EMBEDDER, RETRY & DIMENSION SAFETY TESTS
# =============================================================================

import ssl
from src.rag.embeddings import GeminiEmbedder, EmbeddingError
from src.rag.vector_store import IncompatibleEmbeddingDimensionError


def test_embed_one_chunk_returns_one_embedding():
    """Verify that a single chunk produces exactly one embedding without task_type."""
    embedder = GeminiEmbedder(api_key="mock_key")
    mock_client = MagicMock()
    mock_response = MagicMock()
    mock_item = MagicMock()
    mock_item.values = [0.123] * 768
    mock_response.embeddings = [mock_item]
    mock_client.models.embed_content.return_value = mock_response

    embedder._client = mock_client

    result = embedder.embed_documents(["Single chunk text"])

    assert len(result) == 1
    assert len(result[0]) == 768
    # Verify embed_content was called once with a list of Content objects (not raw strings)
    mock_client.models.embed_content.assert_called_once()
    call_kwargs = mock_client.models.embed_content.call_args[1]
    assert call_kwargs["model"] == "gemini-embedding-2"
    assert len(call_kwargs["contents"]) == 1
    assert hasattr(call_kwargs["contents"][0], "parts")
    # Verify no task_type was passed
    assert "config" not in call_kwargs or call_kwargs["config"] is None


def test_embed_multiple_chunks_returns_same_number_of_embeddings():
    """Verify that embedding 58 chunks produces exactly 58 embeddings batched in groups of 10."""
    embedder = GeminiEmbedder(api_key="mock_key", batch_size=10)
    mock_client = MagicMock()

    # Dynamic return value generating matching number of embeddings per batch
    def fake_embed_content(*args, **kwargs):
        contents = kwargs.get("contents", [])
        mock_resp = MagicMock()
        mock_resp.embeddings = [MagicMock(values=[0.1] * 768) for _ in contents]
        return mock_resp

    mock_client.models.embed_content.side_effect = fake_embed_content
    embedder._client = mock_client

    test_58_chunks = [f"This is chunk number {i} from Pathway PDF." for i in range(58)]
    result = embedder.embed_documents(test_58_chunks)

    assert len(result) == 58
    # With batch_size 10: 58 chunks -> 6 batches (10, 10, 10, 10, 10, 8)
    assert mock_client.models.embed_content.call_count == 6


def test_embed_query_returns_one_embedding():
    """Verify that query embedding produces one vector in the same space without task_type."""
    embedder = GeminiEmbedder(api_key="mock_key")
    mock_client = MagicMock()
    mock_response = MagicMock()
    mock_item = MagicMock()
    mock_item.values = [0.456] * 768
    mock_response.embeddings = [mock_item]
    mock_client.models.embed_content.return_value = mock_response

    embedder._client = mock_client

    result = embedder.embed_query("What are the key requirements?")

    assert isinstance(result, list)
    assert len(result) == 768
    mock_client.models.embed_content.assert_called_once()
    call_kwargs = mock_client.models.embed_content.call_args[1]
    assert call_kwargs["model"] == "gemini-embedding-2"
    assert len(call_kwargs["contents"]) == 1


def test_embed_retry_behavior_on_ssl_error():
    """Verify retry with exponential backoff on SSL EOF error, resetting the client."""
    embedder = GeminiEmbedder(api_key="mock_key", max_retries=3, initial_backoff=0.01)
    mock_client = MagicMock()

    success_resp = MagicMock()
    success_resp.embeddings = [MagicMock(values=[0.5] * 768)]

    # Fail on first attempt with SSL EOF, then succeed on second attempt
    ssl_error = ssl.SSLEOFError("EOF occurred in violation of protocol (_ssl.c:1002)")
    mock_client.models.embed_content.side_effect = [ssl_error, success_resp]

    with patch("src.rag.embeddings.genai.Client", return_value=mock_client):
        result = embedder.embed_documents(["Text to embed"])

    assert len(result) == 1
    assert mock_client.models.embed_content.call_count == 2


def test_embed_retry_exhaustion_raises_embedding_error():
    """Verify that after exhausting retries on SSL EOF, EmbeddingError is raised with clear details."""
    embedder = GeminiEmbedder(api_key="mock_key", max_retries=2, initial_backoff=0.01)
    mock_client = MagicMock()
    ssl_error = ssl.SSLEOFError("EOF occurred in violation of protocol (_ssl.c:1002)")
    mock_client.models.embed_content.side_effect = ssl_error

    with patch("src.rag.embeddings.genai.Client", return_value=mock_client):
        with pytest.raises(EmbeddingError) as exc_info:
            embedder.embed_documents(["Some text"])

    assert "after 2 attempts" in str(exc_info.value)
    assert "EOF occurred in violation of protocol" in str(exc_info.value)


def test_failed_embedding_not_inserted_into_chromadb(tmp_path):
    """Verify that a failed embedding never inserts records into ChromaDB."""
    store = VectorStore(persist_directory=tmp_path / "test_safety", collection_name="safety_col")
    assert store.count() == 0

    embedder = GeminiEmbedder(api_key="mock_key", max_retries=1)
    mock_client = MagicMock()
    mock_client.models.embed_content.side_effect = ConnectionResetError("Connection reset by peer")
    embedder._client = mock_client

    chunks = [{"chunk_id": "c1", "text": "Important text", "source": "doc.pdf", "page": 1, "chunk_index": 0}]

    # Ingestion flow simulation: if embedding fails, add_chunks is never reached
    with pytest.raises(EmbeddingError):
        embeddings = embedder.embed_documents([chunks[0]["text"]])
        store.add_chunks(chunks, embeddings)

    # Database remains pristine
    assert store.count() == 0


def test_incompatible_embedding_dimension_detection(tmp_path):
    """Verify that mixing vectors of different dimensions is detected and prevented."""
    store = VectorStore(persist_directory=tmp_path / "test_dims", collection_name="dims_col")

    chunks1 = [{"chunk_id": "c1", "text": "Text 1", "source": "f1.txt", "page": None, "chunk_index": 0}]
    # 768-dim embeddings from prior model
    embs1 = [[0.1] * 768]
    store.add_chunks(chunks1, embs1)
    assert store.count() == 1

    chunks2 = [{"chunk_id": "c2", "text": "Text 2", "source": "f2.txt", "page": None, "chunk_index": 0}]
    # Incompatible 1536-dim embeddings
    embs2 = [[0.2] * 1536]

    with pytest.raises(IncompatibleEmbeddingDimensionError) as exc_info:
        store.add_chunks(chunks2, embs2)

    assert "768-dimensional" in str(exc_info.value)
    assert "1536-dimensional" in str(exc_info.value)
    assert "Clear Knowledge Base" in str(exc_info.value)


def test_empty_input_embedding():
    """Verify empty document list and empty query handling."""
    embedder = GeminiEmbedder(api_key="mock_key")
    assert embedder.embed_documents([]) == []

    with pytest.raises(ValueError) as exc:
        embedder.embed_query("")
    assert "Query cannot be empty" in str(exc.value)

    with pytest.raises(ValueError):
        embedder.embed_query("   ")


